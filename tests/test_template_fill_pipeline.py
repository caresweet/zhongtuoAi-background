# -*- coding: utf-8 -*-
"""新管线（模板填空 + 缺口补写 + 规则校验）的回归测试。

运行：cd backend && python -m pytest tests/test_template_fill_pipeline.py -v
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest


# ═══════════════════════════════════════════════════════════════════
# 1. 槽位类型学：固定话术豁免 + 类型守卫
# ═══════════════════════════════════════════════════════════════════

def test_fixed_phrase_exempts_regulations():
    from app.services.template_slot_schema import is_fixed_phrase
    # 法规名所在的段落 → 段落里的年份不是槽位
    ctx = "《中华人民共和国土地管理法实施条例》（2021年9月1日起施行）"
    assert is_fixed_phrase("2021年9月1日", ctx) is True
    # 风险等级判定表里的数字是标准阈值，不是项目数据
    assert is_fixed_phrase("200人", "可能引发200人以上群体性事件的为高风险") is True
    assert is_fixed_phrase("50%", "20%以上～50%以下群众反对的为中风险") is True
    # 普通项目数据不是固定话术
    assert is_fixed_phrase("326342平方米", "根据勘测定界报告，总面积326342平方米") is False


def test_value_fits_unit_guard():
    from app.services.template_slot_schema import value_fits_unit
    assert value_fits_unit("100", "群众支持率100%") is True
    assert value_fits_unit("489.51亩", "支持率100%") is False      # % 槽位不得填亩
    assert value_fits_unit("326342", "总面积326342平方米") is True
    assert value_fits_unit("100%", "总面积326342平方米") is False


def test_invalid_slot_value_rejects_placeholder():
    from app.services.template_slot_schema import is_valid_slot_value
    assert is_valid_slot_value("185", "household_count") is True
    assert is_valid_slot_value("【待补充】", "household_count") is False
    assert is_valid_slot_value("", "household_count") is False
    assert is_valid_slot_value(None, "area_m2") is False
    assert is_valid_slot_value("一百八十五", "household_count") is False


# ═══════════════════════════════════════════════════════════════════
# 2. 就地填充：只改文字、保留 run 格式、固定话术不动
# ═══════════════════════════════════════════════════════════════════

def _template_path():
    from app.services.template_chapter_chunker import _find_template_docx
    return _find_template_docx()


def test_docx_fill_replaces_and_preserves_format():
    tpl = _template_path()
    if not tpl:
        pytest.skip("未找到模板 docx")
    from docx import Document
    from app.services.docx_slot_filler import DocxSlotFiller, chapter_span

    doc = Document(tpl)
    s, e = chapter_span(doc, 5)          # 正文第5章 = 模板第六章
    if s < 0:
        pytest.skip("模板中未定位到第5章")

    target = next((p for p in doc.paragraphs[s:e] if "朱坝街道办事处" in p.text), None)
    if target is None:
        pytest.skip("模板第5章不含预期示例值")
    size_before = [r.font.size for r in target.runs]

    entries = [{"anchor": "朱坝街道办事处", "type": "slot", "field": "org_name"}]
    res = DocxSlotFiller().fill_chapter(
        doc, 5, entries, {"org_name": "淮安市洪泽区三河镇人民政府"})

    assert res.filled_count == 1
    assert "淮安市洪泽区三河镇人民政府" in target.text
    assert "朱坝街道办事处" not in target.text
    # run 格式（字号）保持不变
    assert [r.font.size for r in target.runs] == size_before


def test_docx_fill_missing_field_produces_gap():
    tpl = _template_path()
    if not tpl:
        pytest.skip("未找到模板 docx")
    from docx import Document
    from app.services.docx_slot_filler import DocxSlotFiller

    doc = Document(tpl)
    entries = [{"anchor": "朱坝街道办事处", "type": "slot", "field": "org_name"}]
    res = DocxSlotFiller().fill_chapter(doc, 5, entries, {"org_name": "【待补充】"})
    assert res.filled_count == 0
    assert any(g.reason == "missing_field" and g.severity == "error" for g in res.gaps)


def test_docx_fill_leaves_fixed_text_untouched():
    tpl = _template_path()
    if not tpl:
        pytest.skip("未找到模板 docx")
    from docx import Document
    from app.services.docx_slot_filler import DocxSlotFiller

    doc = Document(tpl)
    before = "\n".join(p.text for p in doc.paragraphs)
    entries = [{"anchor": "江苏洪泽经济开发区管理委员会", "type": "slot", "field": "org_name"}]
    DocxSlotFiller().fill_chapter(doc, 5, entries, {"org_name": "某某镇人民政府"})
    after = "\n".join(p.text for p in doc.paragraphs)
    # 固定话术段落必须原样保留
    assert "根据该事项风险因素的识别和风险等级的判断" in after
    assert len(after) <= len(before) + 200


# ═══════════════════════════════════════════════════════════════════
# 3. 规则引擎：阈值 / 异常值 / 缺失显式化 / 内容规则
# ═══════════════════════════════════════════════════════════════════

def test_metric_rules_fire():
    from app.validation.rule_checker import run_checks
    issues = run_checks({"total_samples": "5", "support_rate": "45", "dept_survey_count": "2"})
    ids = {i["rule_id"] for i in issues}
    assert "survey_min_count" in ids          # 5 < 30
    assert "dept_participants_min" in ids     # 2 < 5
    assert "support_rate_abnormal" in ids     # 45 < 60


def test_metric_missing_is_explicit():
    """指标缺失必须显式产出，不能静默通过。"""
    from app.validation.rule_checker import run_checks
    assert run_checks({"project_name": "x"}) == []              # 默认不含缺失条目
    miss = run_checks({"project_name": "x"}, include_missing=True)
    assert any(m.get("reason") == "metric_missing" for m in miss)


def test_outlier_uses_history_or_fallback():
    from app.validation.history_stats import is_outlier
    flag, info = is_outlier("support_rate", 45, min_n=5, fallback=[60, 99.5])
    assert flag is True
    flag2, _ = is_outlier("support_rate", 98, min_n=5, fallback=[60, 99.5])
    assert flag2 is False


def test_content_rules_fire():
    from app.validation.content_rule_checker import check_text
    md = ("本章从合法性角度分析，不再重复罗列。\n"
          "调查户数和人数以实际调查统计为准。\n"
          "【待补充】\n"
          "| 序号 | 面积总计 | 耕地 |\n| --- | --- | --- |\n| 1 | 8284 | 3 |\n")
    ids = {i["rule_id"] for i in check_text(md)}
    assert "no_placeholder_residue" in ids
    assert "no_vague_survey_phrase" in ids
    assert "no_meta_narration" in ids
    assert "table_unit_required" in ids


# ═══════════════════════════════════════════════════════════════════
# 4. 工厂开关分派
# ═══════════════════════════════════════════════════════════════════

def test_factory_dispatch_by_switch(monkeypatch):
    from app.agent.agents import chapters as C
    from app.config import settings

    monkeypatch.setattr(settings, "CHAPTER_PIPELINE", "legacy", raising=False)
    a = C.get_chapter_agent(3)
    assert type(a).__name__ == "Chapter3Agent"

    monkeypatch.setattr(settings, "CHAPTER_PIPELINE", "template_fill", raising=False)
    b = C.get_chapter_agent(3)
    assert type(b).__name__ == "ChapterWriterAgent"
    assert hasattr(b, "run")


def test_ai_gap_writer_only_writes_whitelisted_sections():
    from app.services.ai_gap_writer import AIGapWriter
    w = AIGapWriter()
    assert w.is_writable(10, "（一）决策合法性分析") is True
    assert w.is_writable(13, "（一）风险防范与化解措施") is True
    assert w.is_writable(5, "（三）工作原则") is True   # 工作原则属第5章应急预案
    # 事实性小节不允许 AI 写
    assert w.is_writable(3, "（一）决策名称") is False
    assert w.is_writable(5, "（六）利益相关者诉求") is False

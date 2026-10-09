#!/usr/bin/env python3
"""把复垦范文 docx 抽成「可照搬章节」语料（带占位符的章 markdown）。

**为什么需要它**：复垦的第 2/5/7/10 章是行业通用模板文（编制依据/质量要求与措施/
费用构成/保障措施），跨样本相似度 0.53~0.78，逐字照搬即可；只有**项目专属词**
（项目名、行政区划）与**金额**需要换成本项目的值。这与稳评的参照件语料
（`scripts/build_reference_corpus.py`）是同一条思路，只是源换成了 docx、
体例保持数字编号（样本是「2.1 编制目的」，不是稳评的「（一）编制目的」）。

产物形状与稳评语料**完全一致**，因此运行期直接复用
`app/services/reference_corpus.py` 的零 LLM 渲染：

    data/reference_reports/<ref_id>/meta.json       源信息、照搬章号、占位符词典
    data/reference_reports/<ref_id>/chapters.json    各章 markdown（text 带 {{占位符}}）
    data/reference_reports/<ref_id>/_review_sheet.md 人工校订表（占位符逐个可核对）

用法：
    cd backend
    python3 scripts/build_reclamation_reference.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Dict, List

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from docx import Document  # noqa: E402

OUT_ROOT = BACKEND / "data" / "reference_reports"

REF_ID = "reclamation_jiuye_zx_v1"
SRC_DOCX = ("/Users/mac/Downloads/江苏盐业井神张兴矿区采输卤扩建项目"
            "ZX9、ZX10井施工便道临时用地土地复垦方案.docx")
SRC_TITLE = "江苏盐业井神张兴矿区采输卤扩建项目ZX9、ZX10井施工便道临时用地土地复垦方案"

# 只抽这四章 —— 跨样本相似度最高（2/5/7/10 分别是编制总则、质量要求与措施、
# 投资估算、保障措施）。第 1/3 章项目专属（相似度 0.09），必须由模型生成。
VERBATIM_CHAPTERS = [2, 5, 7, 10]

# ── 项目专属词 → 占位符（**顺序敏感**：长的先替，避免「淮安市淮安区」被「淮安市」截断）──
TOKEN_MAP: List[tuple] = [
    (r"江苏盐业井神张兴矿区采输卤扩建项目\s*ZX9\s*[、,，]\s*ZX10\s*井施工便道", "{{project_name}}"),
    (r"江苏盐业井神张兴矿区采输卤扩建项目", "{{project_name}}"),
    (r"淮安市淮安区", "{{region}}"),
    (r"淮安市", "{{region}}"),
]

# ── 金额类：**数值随项目变，照抄即事实错误** → 换成数值占位符；
#    取值不到时渲染层会**整句删除**（见 reference_corpus._drop_sentences_with_numeric_slots）。
#    句子里同时含多个金额时，整句删除 → 故这一段用「一句一个占位符」的正常替换即可。
NUMERIC_TOKEN_MAP: List[tuple] = [
    (r"静态总投资\s*[\d.]+\s*万元", "静态总投资{{total_static}}万元"),
    (r"动态总投资\s*[\d.]+\s*万元", "动态总投资{{total_dynamic}}万元"),
    (r"工程施工费\s*[\d.]+\s*万元", "工程施工费{{cost_construction}}万元"),
    (r"其他费用\s*[\d.]+\s*万元", "其他费用{{cost_other}}万元"),
    (r"监测与管护费\s*[\d.]+\s*万元", "监测与管护费{{cost_monitor_manage}}万元"),
    (r"预备费\s*[\d.]+\s*万元", "预备费{{cost_reserve}}万元"),
    (r"静态单位面积投资为\s*[\d.]+\s*万元/亩", "静态单位面积投资为{{unit_static}}万元/亩"),
    (r"动态单位面积投资为\s*[\d.]+\s*万元/亩", "动态单位面积投资为{{unit_dynamic}}万元/亩"),
    # 🔴 这五处「合计X万元」**数额各不相同**，必须各给一个槽位 ——
    #    共用一个占位符的话，取到值时会五处填成同一个数（事实错误）。
    (r"(监测费按工程施工费的[\d.]+%收取，)合计\s*[\d.]+\s*万元",
     r"\1合计{{cost_monitor}}万元"),
    (r"(元/（人\*月）计取，)共计\s*[\d.]+\s*万元", r"\1共计{{cost_care}}万元"),
    (r"(与监测管护费之和的[\d.]+%核定，)合计\s*[\d.]+\s*万元",
     r"\1合计{{cost_basic_reserve}}万元"),
    (r"(价差预备费率进行测算，)合计\s*[\d.]+\s*万元", r"\1合计{{cost_price_reserve}}万元"),
    (r"(计提本次风险金，)合计\s*[\d.]+\s*万元", r"\1合计{{cost_risk}}万元"),
]

# 占位符取值说明（写进 meta，供人工校订与运行期核对）
PLACEHOLDER_VALUES = {
    "project_name": "项目名称（取事实卡 project_name/project_title）",
    "region": "行政区划（由 location 推导，如「淮安市淮安区」）",
    "total_static": "静态总投资（万元，计算引擎）",
    "total_dynamic": "动态总投资（万元，计算引擎）",
    "cost_construction": "工程施工费（万元，计算引擎）",
    "cost_other": "其他费用（万元，计算引擎）",
    "cost_monitor_manage": "监测与管护费（万元，计算引擎）",
    "cost_reserve": "预备费（万元，计算引擎）",
    "unit_static": "静态单位面积投资（万元/亩，计算引擎）",
    "unit_dynamic": "动态单位面积投资（万元/亩，计算引擎）",
    "cost_monitor": "监测费合计（万元，计算引擎）",
    "cost_care": "管护费合计（万元，计算引擎）",
    "cost_basic_reserve": "基本预备费合计（万元，计算引擎）",
    "cost_price_reserve": "价差预备费合计（万元，计算引擎）",
    "cost_risk": "风险金合计（万元，计算引擎）",
}

# 数值类占位符：取不到值 → 整句删除（不写 0、不写「无」）
NUMERIC_SLOTS = ["total_static", "total_dynamic", "cost_construction", "cost_other",
                 "cost_monitor_manage", "cost_reserve", "unit_static",
                 "unit_dynamic", "cost_monitor", "cost_care", "cost_basic_reserve",
                 "cost_price_reserve", "cost_risk"]

# 泄漏自检词：渲染结果里若残留这些，说明替换没做干净。
# 🔴 **不能放行政区划词**（如「淮安」）—— 本项目若恰在同一区划，`{{region}}` 就该
#    渲染出「淮安」，那是正确填充而非泄漏；放进去会把正常成稿判成硬错误。
#    只留**项目专属名**（底本项目的名称/编号/专有词）与**底本的金额数字**。
LEAK_WORDS = ["盐业井神", "井神", "张兴", "ZX9", "ZX10", "采输卤",
              "19.85", "21.88", "6.79"]


def _apply(tokens: List[tuple], text: str) -> str:
    for pat, rep in tokens:
        text = re.sub(pat, rep, text)
    return text


def _chapter_no(para) -> int | None:
    """Heading 1 段落 → 章号（'5 土地复垦质量要求与复垦措施' → 5）。"""
    m = re.match(r'^(\d{1,2})\s+\S', para.text.strip())
    return int(m.group(1)) if m else None


def _sec_no(text: str) -> str | None:
    m = re.match(r'^(\d{1,2}(?:\.\d{1,2}){1,2})\s*\S', text.strip())
    return m.group(1) if m else None


def collect(path: str) -> Dict[int, dict]:
    """按 Heading 层级抽 VERBATIM_CHAPTERS 各章的 markdown。

    docx 用的是真标题样式，比 PDF 的「按字号猜标题」可靠得多：
      Heading 1 → 章；Heading 2 → ## X.Y；Heading 3 → ### X.Y.Z；
      Heading 6 → 表标题（`表N-M …`）；其余 → 正文段。
    """
    doc = Document(path)
    out: Dict[int, dict] = {}
    cur: int | None = None
    buf: List[str] = []

    def flush():
        if cur is not None and buf:
            out[cur] = {"lines": list(buf)}

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = p.style.name
        if style == "Heading 1":
            flush()
            buf = []
            n = _chapter_no(p)
            cur = n if (n in VERBATIM_CHAPTERS) else None
            if cur is not None:
                buf.append(f"# {t}")          # 章标题（渲染层不用，留作分章边界）
            continue
        if cur is None:
            continue

        if style == "Heading 2":
            sec = _sec_no(t)
            buf.append(f"## {sec} {t[len(sec):].strip()}" if sec else f"## {t}")
        elif style == "Heading 3":
            sec = _sec_no(t)
            buf.append(f"### {sec} {t[len(sec):].strip()}" if sec else f"### {t}")
        elif style == "Heading 6":
            buf.append(f"**{t}**")            # 表标题：留作正文加粗，表体本身不在 docx 文本层
        else:
            buf.append(t)
    flush()
    return out


def to_text(lines: List[str]) -> str:
    """行列表 → 段落 markdown（去章标题行，空行分段）。"""
    body = [ln for ln in lines if not ln.startswith("# ")]
    md = "\n\n".join(body)
    md = re.sub(r'\n{3,}', '\n\n', md)
    return md.strip()


def main() -> int:
    src = Path(SRC_DOCX)
    if not src.exists():
        print(f"❌ 源文件不存在：{src}")
        return 2

    print(f"读取 {src.name} …")
    raw = collect(str(src))
    missing = [n for n in VERBATIM_CHAPTERS if n not in raw]
    if missing:
        print(f"❌ 没抽到章：{missing}（检查 Heading 样式或源文件）")
        return 3

    chapters: Dict[str, dict] = {}
    for n in VERBATIM_CHAPTERS:
        md = to_text(raw[n]["lines"])
        md = _apply(TOKEN_MAP, md)
        md = _apply(NUMERIC_TOKEN_MAP, md)
        title = raw[n]["lines"][0][2:].strip()
        title = re.sub(r'^\d+\s+', '', title)
        chapters[str(n)] = {
            "title": title,
            "text": md,
            "chars": len(md),
            "placeholders": sorted({m for m in re.findall(r'\{\{(\w+)\}\}', md)}),
            "source_kind": "docx",
        }
        ph = re.findall(r'\{\{(\w+)\}\}', md)
        print(f"  第{n}章「{title}」{len(md)} 字，占位符 {len(ph)} 处")

    out_dir = OUT_ROOT / REF_ID
    out_dir.mkdir(parents=True, exist_ok=True)

    (out_dir / "chapters.json").write_text(json.dumps({
        "source": SRC_TITLE,
        "ref_id": REF_ID,
        "chapters": chapters,
    }, ensure_ascii=False, indent=1), encoding="utf-8")

    all_ph: Dict[str, int] = {}
    for c in chapters.values():
        for m in re.findall(r'\{\{(\w+)\}\}', c["text"]):
            all_ph[f"{{{{{m}}}}}"] = all_ph.get(f"{{{{{m}}}}}", 0) + 1

    (out_dir / "meta.json").write_text(json.dumps({
        "ref_id": REF_ID,
        "domain": "reclamation",
        "source_docx": SRC_DOCX,
        "source_title": SRC_TITLE,
        "verbatim_chapters": VERBATIM_CHAPTERS,
        # 🔴 体例：样本用「2.1 编制目的」的数字编号，**不做**稳评的「（一）」转换
        "heading_mode": "numeric",
        # 🔴 第2章不做「（三）相关文件」本项目化 —— 那是稳评的编制依据体例，复垦没有
        "chapter2_documents_block": False,
        # 🔴 数值槽位取不到值时整句会被删；若某小节正文只有那一句，就只剩光标题 →
        #    开启「空标题清理」，把「7.1.3 投资情况」这类删不剩正文的小节标题一并去掉。
        "drop_empty_sections": True,
        "placeholder_values": PLACEHOLDER_VALUES,
        "numeric_slots": NUMERIC_SLOTS,
        "leak_words": LEAK_WORDS,
        # 复垦的这些章是行业通用模板文，不含风电式「项目类型专属条目」→ 不过滤
        "project_type_deny": [],
        "placeholders": all_ph,
        "reviewed": False,
    }, ensure_ascii=False, indent=1), encoding="utf-8")

    sheet = ["# 复垦照搬语料 人工校订表", "",
             f"源：{SRC_TITLE}", f"照搬章：{VERBATIM_CHAPTERS}", "",
             "## 占位符清单", ""]
    for k, v in PLACEHOLDER_VALUES.items():
        sheet.append(f"- `{{{{{k}}}}}` —— {v}（出现 {all_ph.get('{{'+k+'}}', 0)} 次）")
    for n in VERBATIM_CHAPTERS:
        sheet += ["", f"## 第{n}章 {chapters[str(n)]['title']}", "",
                  "```markdown", chapters[str(n)]["text"], "```"]
    (out_dir / "_review_sheet.md").write_text("\n".join(sheet), encoding="utf-8")

    print(f"\n✅ 已写入 {out_dir}")
    print(f"   chapters.json / meta.json / _review_sheet.md")
    print(f"   占位符：{all_ph}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

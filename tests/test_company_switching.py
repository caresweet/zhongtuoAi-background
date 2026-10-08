"""公司信息「内外有别」—— 外部用户的报告里**绝不能出现江苏众拓**。

🔴 为什么必须有这个文件
    报告的**公司简介四段、编写单位、封面实施单位**原先全部**写死在装配器里**
    （`report_assembler._add_company_intro` 的字面量循环、
    `_render_personnel_signature_table` 的「编写单位：江苏众拓…」）。
    商用后外部用户出报告，等于**把别人的客户抬头印成我们的**。

    现在这些从 `companies` 表按归属取；取不到显示「待填写」。
    这一组用例直接对**产出的 docx 正文**做断言 —— 光看代码容易漏掉某处兜底。
"""

import re

from docx import Document

from app.services.internal_company import INTERNAL_COMPANY_INTRO, INTERNAL_COMPANY_NAME


def _assemble_doc(company: dict):
    """用最小 state 跑一遍装配，返回 Document。"""
    from pathlib import Path

    from app.services.report_assembler import ReportAssembler

    state = {
        "session_id": "company-switch-test",
        "chapters": {},
        "filled_data": {},
        "_company": company,
        "_company_id": company.get("id"),
        "_domain": "stability",
    }
    path = ReportAssembler().assemble(state)

    # 🔴 `assemble()` 返回的是**相对路径**（如 `generated/xxx.docx`），
    #    直接 `Document(path)` 会 PackageNotFound —— 真实落点在 storage/ 下。
    #    脚本侧的 `gen_full_report.py` 也得做同样的一段解析，属于既有口径。
    p = Path(path)
    if not p.is_absolute():
        base = Path(__file__).resolve().parent.parent
        for cand in (base / p, base / "storage" / p):
            if cand.exists():
                p = cand
                break
    return Document(str(p))


def _assemble_with_company(company: dict) -> str:
    """用最小 state 跑一遍装配，返回全文（段 + 表格）。"""
    doc = _assemble_doc(company)
    parts = [pp.text for pp in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_internal_company_report_contains_our_intro():
    """内部公司（江苏众拓）→ 简介与编写单位都用我们的。"""
    full = _assemble_with_company({
        "id": 1, "name": INTERNAL_COMPANY_NAME, "is_internal": True,
        "intro_text": INTERNAL_COMPANY_INTRO,
        "implement_unit": INTERNAL_COMPANY_NAME,
    })
    assert INTERNAL_COMPANY_NAME in full, "内部公司的报告里应出现我们的公司名"
    # 简介首段的前十几个字应当出现（证明用的是库里的 intro_text 而不是别处）
    assert INTERNAL_COMPANY_INTRO[:16] in full, "应渲染库中的公司简介"


def test_external_user_report_never_leaks_our_company():
    """**核心断言**：没有公司信息时，报告里一个字都不能提江苏众拓。"""
    full = _assemble_with_company({})
    assert "江苏众拓" not in full, (
        "外部用户的报告里出现了我们的公司名 —— 这正是「写死兜底」会造成的后果。\n"
        "命中片段：" + str([l for l in full.split("\n") if "江苏众拓" in l][:5])
    )
    # 并应给出「待填写」让用户知道要补
    assert "待填写" in full, "公司信息缺失时应显示「待填写」占位"


def test_external_user_report_uses_their_own_company():
    """外部用户填了自己的公司 → 报告用**他的**，且仍然不提我们。"""
    full = _assemble_with_company({
        "id": 2, "name": "某某工程咨询有限公司", "is_internal": False,
        "intro_text": "某某工程咨询有限公司成立于 2010 年，主营社会稳定风险评估。",
        "implement_unit": "某某工程咨询有限公司",
    })
    assert "某某工程咨询有限公司" in full
    assert "江苏众拓" not in full
    assert full.count("某某工程咨询有限公司") >= 2, "简介与编写单位都该用他的公司名"


def test_migration_does_not_assign_company_to_new_users():
    """迁移**只能**给我们自己的引导账号挂公司，不能顺手把别人的也挂上。

    🔴 实测踩过：`_backfill_user_company` 写成「把所有 company_id 为空的账号
       都挂到内部公司」，于是**刚建好的、还没填公司信息的外部账号**被自动挂成
       江苏众拓 —— 它的报告会用我们的简介/资质，人员名单也跟着串成我们的 22 人。
       「内外有别」当场失效，而且不报任何错。
    """
    from sqlalchemy import create_engine

    from app.database.auth_db import _promote_bootstrap_admin

    # 🔴 必须用 SQLAlchemy 的 Connection：迁移里用的是 `text()` 语句，
    #    原生的 sqlite3.Connection 接不了。
    engine = create_engine("sqlite://")
    with engine.begin() as conn:
        # 🔴 `exec_driver_sql` 一次只能跑一条语句（sqlite3 的限制），逐条建
        conn.exec_driver_sql(
            "CREATE TABLE companies (id INTEGER PRIMARY KEY, name TEXT, is_internal INTEGER)")
        conn.exec_driver_sql(
            "CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT, role TEXT, "
            "is_active INTEGER, is_superuser INTEGER, company_id INTEGER)")
        conn.exec_driver_sql(
            "INSERT INTO companies (id, name, is_internal) "
            "VALUES (1, '江苏众拓项目代理咨询有限公司', 1)")
        conn.exec_driver_sql("INSERT INTO users VALUES (1, 'admin', 'user', 1, 0, NULL)")
        conn.exec_driver_sql("INSERT INTO users VALUES (2, 'waibu01', 'user', 1, 0, NULL)")
        conn.exec_driver_sql("INSERT INTO users VALUES (3, 'waibu02', 'user', 1, 0, 2)")
        _promote_bootstrap_admin(conn)

        rows = dict(conn.exec_driver_sql("SELECT username, company_id FROM users").fetchall())
        assert rows["admin"] == 1, "引导账号应挂到内部公司"
        assert rows["waibu01"] is None, f"外部账号被误挂成公司 {rows['waibu01']} —— 内外有别失效"
        assert rows["waibu02"] == 2, "已有公司的账号不该被改动"

        admin_role = conn.exec_driver_sql(
            "SELECT role, is_superuser FROM users WHERE username='admin'").fetchone()
        assert tuple(admin_role) == ("admin", 1), "引导账号应被提权"


def test_external_report_carries_no_our_hardcoded_assets():
    """**图片层面的泄露**：外部用户的报告里不能贴我们的营业执照/资质图。

    🔴 这批图（`storage/extracted_imgs/stability_cert_*`）是**固定文件**、
       代码里按文件名约定直接扫进来，与文本无关 —— 光断言文本抓不到。
       实测确实会把我们的营业执照贴进外部用户的报告。
    """
    external = _assemble_doc({})
    internal = _assemble_doc({
        "id": 1, "name": INTERNAL_COMPANY_NAME, "is_internal": True,
        "intro_text": INTERNAL_COMPANY_INTRO,
        "implement_unit": INTERNAL_COMPANY_NAME,
    })
    n_ext = len(external.inline_shapes)
    n_int = len(internal.inline_shapes)
    assert n_int > 0, "内部公司的报告应当带得上资质图（否则这条断言失去意义）"
    assert n_ext == 0 or n_ext < n_int, (
        f"外部用户的报告里出现了 {n_ext} 张图（内部公司为 {n_int} 张）——"
        f" 很可能把我们的营业执照/资质图贴出去了"
    )
    # 文本上也要给出去补资料的提示
    txt = "\n".join(p.text for p in external.paragraphs)
    assert "待填写" in txt

"""组装期的**表格完整性**回归。

三处真实缺陷（2026-09-23 实测稿「洪拟征告〔2026〕10号…_0923_1112.docx」）：

① 章号解析表只写到「十」→ 第 11 章及以后一律回落到 0，
   表号重编与同章表格去重**整体静默失效**，第12章因此留下两张一样的测评表。
② 正文里「第11章识别的主要风险因素，经问卷调查…」也被当成章标题，
   章号被带偏到「第11章」。
③ 相邻两张 markdown 表之间没有正文/空行时被并成**一张** docx 表 ——
   第12章「措施前评分表 + 措施前后得分对比表」合成了一张 29 行的怪表，
   表中间冒出「序号|评估维度|措施前得分|措施后得分」这行新表头。
"""

import re

from docx import Document


# ════════════════════════════════════════════════════════
# ① 章号解析（含第 11 章及以后）
# ════════════════════════════════════════════════════════

def test_cn_num_to_int_covers_teens_and_beyond():
    from app.agent.chapter_definitions import cn_num_to_int

    cases = {"一": 1, "十": 10, "十一": 11, "十二": 12, "十四": 14,
             "二十": 20, "二十三": 23, "一百零五": 105,
             "7": 7, "14": 14}
    for s, want in cases.items():
        assert cn_num_to_int(s) == want, f"{s} → {cn_num_to_int(s)}，应为 {want}"


def test_chapter_no_of_heading_rejects_body_sentence():
    """正文里提到「第11章…」不是章标题，不得被当成章标题。"""
    from app.agent.chapter_definitions import chapter_no_of_heading as ch

    assert ch("第十一章、分析预测风险源") == 11
    assert ch("第十四章、附件") == 14
    assert ch("第十二章 确定风险等级") == 12
    # 🔴 实测的误匹配句（以「第11章」开头）
    assert ch("第11章识别的主要风险因素，经问卷调查、部门意见征询和实地核查，"
              "发生概率较小，影响程度较低，初始风险等级均评定为低风险。") == 0
    assert ch("表12-1 直接定性指标") == 0
    assert ch("（三）主要风险因素初始风险等级表") == 0


# ════════════════════════════════════════════════════════
# ③ 相邻两张 markdown 表必须拆成两张
# ════════════════════════════════════════════════════════

_TWO_TABLES_MD = """### （二）稳评风险等级确定量化指标评价体系

量化指标评价体系见下表。

| 测评指标 | 权重 | 测评项目 | 评分 |
|---|---|---|---|
| 合法性 | 10 | 事项实施主体是否符合规定 | 3 |
| 合计 | 10 | 小计 | 3 |
| 序号 | 评估维度 | 措施前得分 | 措施后得分 |
|---|---|---|---|
| 1 | 合法性 | 0 | 0 |
| 2 | 合计 | 0 | 0 |
"""


def _fresh_assembler():
    from app.services.report_assembler import ReportAssembler

    a = ReportAssembler()
    a._inserted_images = set()
    a._chapter_image_counter = {}
    a._system_rendered_tables = set()
    a._rendered_extracted_tbls = set()
    a._session_state = {}
    return a


def test_back_to_back_markdown_tables_are_not_merged():
    a = _fresh_assembler()
    d = Document()
    a._setup_page(d)
    a._add_chapter(d, 12, _TWO_TABLES_MD, {}, {}, [], {})

    assert len(d.tables) == 2, (
        f"相邻两表被并成了 {len(d.tables)} 张；表头依次为 "
        f"{[c.text.strip() for c in d.tables[0].rows[0].cells]}"
    )
    first = [c.text.strip() for c in d.tables[0].rows[0].cells]
    second = [c.text.strip() for c in d.tables[1].rows[0].cells]
    assert first[:2] == ["测评指标", "权重"]
    assert second[:2] == ["序号", "评估维度"], f"第二张表表头错了：{second}"
    assert len(d.tables[0].rows) == 3, "第一张表应只有表头 + 2 行数据"
    assert len(d.tables[1].rows) == 3, "第二张表应只有表头 + 2 行数据"


def test_chapter12_scoring_table_keeps_all_six_columns():
    """第12章的量化评分表必须是 6 列 —— 旧代码按章号硬裁成 4 列，
    把「评分标准」「得分」两列切掉了，与提示词要求直接冲突。"""
    a = _fresh_assembler()
    d = Document()
    a._setup_page(d)
    a._add_chapter(d, 12, """### （二）稳评风险等级确定量化指标评价体系

量化指标评价体系见下表。

| 测评指标 | 权重 | 测评项目 | 评分 | 评分标准 | 得分 |
|---|---|---|---|---|---|
| 合法性 | 10 | 事项实施主体是否符合规定 | 0 | 全部符合不计分，不符合计10分 | 0 |
| 合计 | 100 | 小计 | 0 |  | 0 |
""", {}, {}, [], {})

    assert len(d.tables) == 1
    hdr = [c.text.strip() for c in d.tables[0].rows[0].cells]
    assert hdr == ["测评指标", "权重", "测评项目", "评分", "评分标准", "得分"], (
        f"第12章评分表列被裁掉了：{hdr}"
    )


# ════════════════════════════════════════════════════════
# ① 同章重复表去重（第 11 章及以后也要生效）
# ════════════════════════════════════════════════════════

def _add_table(doc, headers, data_rows):
    t = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for ri, row in enumerate(data_rows, start=1):
        for ci, v in enumerate(row):
            t.rows[ri].cells[ci].text = v
    return t


def test_dedup_duplicate_tables_works_past_chapter_ten():
    """第十二章（>十）里两张一样的测评表必须去重，且连孤立表题一起删。

    🔴 真实稿里两张表**没有挨着** —— 中间夹着正文
       「第11章识别的主要风险因素，经问卷调查…」。旧写法把这句当成了章标题，
       把第二张表划到「第11章」，于是同章判定失效、重复表留了下来。
       这里就按真实顺序摆：表 → 那句正文 → 重复表。
    """
    a = _fresh_assembler()
    d = Document()
    d.add_paragraph("第十二章、确定风险等级")
    headers = ["测评指标", "权重", "测评项目", "评分"]
    rows = [["合法性", "10", "事项实施主体是否符合规定", "3"],
            ["合计", "10", "小计", "3"]]

    d.add_paragraph("表12-2 稳评风险等级确定量化指标评价体系（措施前）")
    _add_table(d, headers, rows)
    d.add_paragraph("第11章识别的主要风险因素，经问卷调查、部门意见征询和实地核查，"
                    "发生概率较小，影响程度较低，初始风险等级均评定为低风险。")
    d.add_paragraph("表11-1 稳评风险等级确定量化指标评价体系（措施前）")
    _add_table(d, headers, rows)

    assert len(d.tables) == 2
    a._dedup_duplicate_tables(d)

    assert len(d.tables) == 1, "同章重复的测评表没被去掉"
    titles = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith("表")]
    assert titles == ["表12-2 稳评风险等级确定量化指标评价体系（措施前）"], (
        f"被删表的孤立表题没清掉：{titles}"
    )


def test_dedup_keeps_same_table_in_different_chapters():
    """同表头但**不同章**的表不得互删（措施前/措施后分别在第12、4章）。"""
    a = _fresh_assembler()
    d = Document()
    headers = ["测评指标", "权重", "测评项目", "评分"]
    rows = [["合法性", "10", "事项实施主体是否符合规定", "3"]]

    d.add_paragraph("第四章、社会稳定风险评估工作报告")
    _add_table(d, headers, rows)
    d.add_paragraph("第十二章、确定风险等级")
    _add_table(d, headers, rows)

    a._dedup_duplicate_tables(d)
    assert len(d.tables) == 2, "跨章的同名表被误删了"


def test_table_title_renumbering_works_past_chapter_ten():
    """第12章的表号必须重编成 表12-N，而不是保留 LLM 写错的「表11-1」。"""
    a = _fresh_assembler()
    d = Document()
    d.add_paragraph("第十二章、确定风险等级")
    d.add_paragraph("表11-1 稳评风险等级确定量化指标评价体系（措施前）")
    _add_table(d, ["a", "b"], [["1", "2"]])
    d.add_paragraph("表3-7 群众意见分析")
    _add_table(d, ["c", "d"], [["3", "4"]])

    a._dedup_table_titles(d)
    titles = [p.text.strip() for p in d.paragraphs if p.text.strip().startswith("表")]
    assert titles == ["表12-1 稳评风险等级确定量化指标评价体系（措施前）", "表12-2 群众意见分析"], titles
    assert not any(re.search(r"表11-1", p.text) for p in d.paragraphs)


# ════════════════════════════════════════════════════════
# ④ 系统标记绝不能印进正文
# ════════════════════════════════════════════════════════

def test_unrendered_system_markers_never_reach_the_document():
    """`[TABLE:…]` / `[SYS_TABLE:…]` 是给装配器看的，漏进交付稿就是天书。

    实测 2026-10-28：业主在成稿里看到「[TABLE:ch9_public_survey]」字样。
    """
    from docx import Document

    a = _fresh_assembler()
    a._domain = "stability"
    a._filled = {}
    d = Document()
    a._setup_page(d)
    md = ("### （一）问卷调查情况\n\n引导语。\n\n"
          "[TABLE:ch9_unknown_table]\n\n"
          "[SYS_TABLE:9-9]\n\n"
          "[TABLE: ch9_public_survey]\n\n"
          "解读句。\n")
    a._add_chapter(d, 9, md, {}, {}, [], {})
    text = "\n".join(p.text for p in d.paragraphs)
    assert "[TABLE" not in text and "[SYS_TABLE" not in text
    assert "引导语" in text and "解读句" in text


# ════════════════════════════════════════════════════════
# ⑤ 稳评总分跨章一致
# ════════════════════════════════════════════════════════

def test_risk_score_is_unified_across_chapters():
    """第13章从范文抄来的「综合风险指数 12.5 分」必须对齐第12章的 W=0 分。"""
    from app.validation.content_guardrails import fix_risk_score_consistency

    chapters = {
        12: {"markdown": "扣分合计即为稳评总分W。\n\n综上，措施前稳评总分 W=0 分，判定为低风险。"},
        13: {"markdown": "第11章识别出本项目六项风险因素，经第12章量化评分，"
                         "措施前综合风险指数为12.5分，风险等级为低风险。"},
    }
    out, fixes = fix_risk_score_consistency(chapters)
    assert "12.5" not in out[13]["markdown"]
    assert "综合风险指数为0分" in out[13]["markdown"]
    assert fixes and fixes[0][0] == 13


def test_risk_score_fix_keeps_consistent_text_untouched():
    from app.validation.content_guardrails import fix_risk_score_consistency

    md12 = "措施前稳评总分 W=0 分。"
    chapters = {12: {"markdown": md12}, 13: {"markdown": "综合风险指数为0分，属低风险。"}}
    out, fixes = fix_risk_score_consistency(chapters)
    assert fixes == [] and out[13]["markdown"] == chapters[13]["markdown"]


# ════════════════════════════════════════════════════════
# ⑥ 评审表摘录不得带系统标记
# ════════════════════════════════════════════════════════

def test_review_table_extracts_clean_text():
    """评审表「稳评方法与过程」一格曾印出 `[TABLE:ch9_public_survey]`（2026-10-28 实测）。"""
    from app.services.review_table_service import _chapter_md, _extract_from_chapters

    chapters = {
        9: {"markdown": "（一）问卷调查情况\n\n引导语。\n\n[TABLE:ch9_public_survey]\n\n解读句。"},
        6: {"markdown": "![现场图](images/x.jpg)\n\n第6章正文。"},
    }
    txt = _extract_from_chapters(chapters)
    assert "[TABLE" not in txt and "![" not in txt
    assert "引导语" in txt and "解读句" in txt
    assert "[TABLE" not in _chapter_md(chapters, 9)


def test_review_table_drops_dangling_table_leadin():
    """评审表里不能留「…如下表所示：」这种悬空句子（表格并未复制进来）。"""
    from app.services.review_table_service import _chapter_md

    md = ("个人问卷统计情况如下表所示：\n\n"
          "本次调查共发放个人问卷50份，回收有效问卷50份，回收率100.0%。\n\n"
          "[TABLE:ch9_public_survey]\n")
    out = _chapter_md({9: {"markdown": md}}, 9)
    assert "如下表所示" not in out
    assert "50份" in out and "100.0%" in out


def test_chapter_self_reference_is_stripped_but_crossref_kept():
    """正文/评审表里不写「第11章识别出…经第12章量化评分」；合法的「详见第12章表12-4」保留。"""
    from app.validation.content_guardrails import strip_chapter_self_reference

    chapters = {
        13: {"markdown": "第11章识别出本项目六项风险因素，经第12章量化评分，"
                         "措施前综合风险指数为0分，风险等级为低风险。"},
        11: {"markdown": "…初始风险等级汇总详见第12章表12-4。"},
    }
    out, n = strip_chapter_self_reference(chapters)
    assert n == 2
    assert "第11章" not in out[13]["markdown"] and "第12章" not in out[13]["markdown"]
    assert "识别出本项目六项风险因素" in out[13]["markdown"]
    assert "详见第12章表12-4" in out[11]["markdown"]


# ════════════════════════════════════════════════════════
# ⑦ 评审表各栏的取数位置要对上参照件（洞庭湖路）
# ════════════════════════════════════════════════════════

def test_review_stakeholders_is_group_list_not_survey_stats():
    """「利益相关者情况」要写各类群体清单（参照件写法），不是问卷统计。"""
    from app.services.review_table_service import _build_review_data

    filled = {"org_name": "淮安市洪泽区人民政府",
              "location": "朱坝街道办事处、三圩社区、二组、三组、六组",
              "household_count": "50", "total_samples": "50", "support_rate": "100%"}
    d = _build_review_data(filled, {}, {})
    s = d["stakeholders"]
    assert "稳评责任单位（淮安市洪泽区人民政府）" in s
    assert "街道（朱坝街道）" in s
    assert "集体经济组织（三圩社区）" in s
    assert "被征地农户" in s and "50" in s


def test_review_risk_factors_are_numbered_subheads():
    """「主要风险因素」是 1./2. 的风险名称清单（参照件写法）。"""
    from app.services.review_table_service import _build_review_data

    md = ("## 11 分析预测风险源\n\n### （一）征收程序方面的风险\n正文甲。\n\n"
          "### （二）土地征收补偿方案方面的风险\n正文乙。\n")
    d = _build_review_data({}, {11: {"markdown": md}}, {})
    assert d["risk_factors"] == "1.征收程序方面的风险\n2.土地征收补偿方案方面的风险"


def test_review_risk_measures_carry_title_and_first_para():
    from app.services.review_table_service import _build_review_data

    md = ("## 13 防范和化解风险的工作建议\n\n### （一）规范征地程序\n严格履行相关程序。\n\n"
          "### （二）合法制定补偿方案\n做好政策解释。\n")
    d = _build_review_data({}, {13: {"markdown": md}}, {})
    assert d["risk_measures"].startswith("1.规范征地程序\n严格履行相关程序。")
    assert "2.合法制定补偿方案" in d["risk_measures"]


# ════════════════════════════════════════════════════════
# ⑧ 去重 & 表9-1 取数
# ════════════════════════════════════════════════════════

def test_dedup_removes_repeated_fragment_and_paragraph():
    from app.validation.content_guardrails import dedup_repeated_text

    md = "本次拟征收土地面积489.5亩，涉及户数50户。50户。补偿方案制定应严格对照标准。"
    out, n = dedup_repeated_text(md)
    assert n == 1 and "50户。50户" not in out and "涉及户数50户。补偿方案" in out

    out2, n2 = dedup_repeated_text("第一段内容。\n\n第一段内容。\n\n第二段。")
    assert n2 == 1 and out2.count("第一段内容") == 1


def test_same_table_number_is_rendered_once():
    """系统已渲染表6-1后，模型手写的同名表必须丢弃（实测重复且错列）。"""
    a = _fresh_assembler()
    a._rendered_captions = set()
    d = Document()
    a._setup_page(d)
    lines = ["| 单位 | 面积 |", "|---|---|", "| 甲 | 100 |"]
    a._render_md_table(d, lines, 6, {}, caption="表6-1 土地分类面积表")
    a._render_md_table(d, lines, 6, {}, caption="表6-1 土地分类面积表")
    assert len(d.tables) == 1


def test_survey_table_reads_dimension_text_instead_of_zeros():
    """表9-1 必须用 `survey_dimension_text` 的逐题清点，不能印全 0。"""
    from app.services.table_generation_service import TableGenerationService

    state = {"filled_data": {
        "survey_total_count": 50,
        "survey_dimension_text": (
            "【问卷调查分项统计（共 50 份，逐题勾选结果）】\n"
            "  · 您是：本地居民（48人）、租住本区（1人）、附近上班（1人）\n"
            "  · 您的年龄：16～35（5人）、36～55（26人）、56 以上（19人）\n"
            "  · 对该决策是否支持：支持（50人）")}}
    t = TableGenerationService.generate_table_3_1(state)
    flat = {tuple(r[:2]): r[2] for r in t["rows"]}
    assert flat[("您是", "本地居民")] == "48"
    assert flat[("您的年龄", "36～55")] == "26"
    assert flat[("对该决策是否支持", "支持")] == "50"


def test_risk_level_table_rebuilt_to_match_ch11():
    """表12-4 的风险条目必须与第11章小节一一对应（实测 6 项被写成 5 条自拟名）。"""
    from app.validation.content_guardrails import rebuild_risk_level_table

    ch11 = ("## 11 分析预测风险源\n\n### （一）甲风险\n正文。该事项风险程度为一般。\n\n"
            "### （二）乙风险\n正文。该事项风险程度为较小。\n\n"
            "### （三）丙风险\n正文。该事项风险程度为一般。\n")
    ch12 = ("## 12 确定风险等级\n\n| 序号 | 风险因素 | 发生概率 | 影响程度 | 初始风险等级 |\n"
            "| --- | --- | --- | --- | --- |\n| 1 | 模型自拟名 | 低 | 一般 | 低风险 |\n\n结论段。\n")
    chapters = {11: {"markdown": ch11}, 12: {"markdown": ch12}}
    assert rebuild_risk_level_table(chapters) == 1
    md = chapters[12]["markdown"]
    assert "| 1 | 甲风险 | 一般 |" in md
    assert "| 2 | 乙风险 | 较小 |" in md
    assert "模型自拟名" not in md and "结论段。" in md

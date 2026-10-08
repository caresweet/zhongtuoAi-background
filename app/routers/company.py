"""公司档案与人员库 —— `/api/v1/company/*`（**登录即可**，不是管理员专属）。

🔴 为什么从 knowledge 路由里搬出来
    人员选择是**每个用户都要用**的能力（「选择报告人员」多选框 → 报告前置的
    工作组名单与培训证书取图）。而知识库按业主要求改成**管理员专属**，
    再放在一起，普通用户就选不了人员了。

🔴 与「内外有别」的关系
    这里的每条数据都带 `company_id`：公司和人员都按**当前用户的公司**过滤。
    外部用户没有公司时返回空列表（而不是看到我们的名册）。
"""

from __future__ import annotations

import json
import logging
from typing import List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from sqlalchemy import select, update

from app.database.auth_db import get_auth_db
from app.deps import get_current_user, is_admin
from app.models.user import Company, User
from app.schemas.knowledge import ApiResponse

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1/company", tags=["公司档案"])


# ═══════════════════════════════════════════════════════════════
# 公司档案
# ═══════════════════════════════════════════════════════════════

def _company_dict(c: Optional[Company]) -> dict:
    if c is None:
        return {}
    return {
        "id": c.id, "name": c.name or "", "is_internal": bool(c.is_internal),
        "short_name": c.short_name or "", "intro_text": c.intro_text or "",
        "implement_unit": c.implement_unit or "",
        "business_license_no": c.business_license_no or "",
        "qualification_text": c.qualification_text or "",
        "address": c.address or "", "contact": c.contact or "",
        # 🔴 负责人 / 职务（评审表「稳评实施单位·负责人 / 职务」两格）
        "leader_name": c.leader_name or "", "leader_title": c.leader_title or "",
    }


async def _load_company(user: User, db) -> Optional[Company]:
    if not user.company_id:
        return None
    return (await db.execute(
        select(Company).where(Company.id == user.company_id)
    )).scalar_one_or_none()


@router.get("/profile", response_model=ApiResponse)
async def get_profile(
    user: User = Depends(get_current_user),
    db=Depends(get_auth_db),
):
    """当前用户公司的档案。**没填过时返回空字段** —— 报告里会显示「待填写」。"""
    c = await _load_company(user, db)
    return ApiResponse(
        data={"company": _company_dict(c), "has_company": c is not None},
        message="" if c else "尚未填写公司信息，报告中相关位置将显示「待填写」",
    )


@router.put("/profile", response_model=ApiResponse)
async def update_profile(
    payload: dict,
    user: User = Depends(get_current_user),
    db=Depends(get_auth_db),
):
    """新建/更新本公司档案。

    🔴 外部用户**不能**改到我们的内部公司记录：`is_internal=True` 的行
       只有管理员能改。
    """
    c = await _load_company(user, db)

    if c is None:
        # 首次填写 → 建公司并挂到当前账号
        name = str(payload.get("name") or "").strip()
        if not name:
            raise HTTPException(status_code=400, detail="请填写公司名称")
        c = Company(name=name, is_internal=False)
        db.add(c)
        await db.flush()
        user.company_id = c.id
    elif c.is_internal and not is_admin(user):
        raise HTTPException(status_code=403, detail="内部公司资料仅管理员可修改")

    # 只允许改白名单字段（其余字段由管理员后台维护）
    for key in ("name", "short_name", "intro_text", "implement_unit",
                "business_license_no", "qualification_text", "address", "contact",
                # 🔴 负责人 / 职务：评审表实施单位那两格按公司取，
                #    所以必须能在「公司信息」页维护（每家公司的负责人不同）
                "leader_name", "leader_title"):
        if key in payload and payload[key] is not None:
            setattr(c, key, str(payload[key]).strip())

    await db.flush()
    return ApiResponse(message="已保存", data={"company": _company_dict(c)})


@router.get("/issuer-candidates", response_model=ApiResponse)
async def issuer_candidates(
    user: User = Depends(get_current_user),
    db=Depends(get_auth_db),
):
    """报告可选的「出具单位」候选。

    - 内部账号：我司各机构（companies.is_internal = 1 的全部）
    - 外部账号：本账号关联的那家公司
    - 账号未关联公司：返回空列表，**不报错**（生成流程不得被打断）

    🔴 为什么不用 `company_service.internal_company_id()`
       那个函数取的是 **id 最小** 的那一家内部公司（单公司时代的写法），
       据此判断「是否内部」会让测绘公司下的用户被当成外部。这里改判
       **用户自己所属公司** 的 `is_internal` 标志 —— 内部用户才能看到两家。
    """
    me = await _load_company(user, db)
    if me is not None and me.is_internal:
        rows = (await db.execute(
            select(Company).where(Company.is_internal == True)  # noqa: E712
            .order_by(Company.id)
        )).scalars().all()
    elif me is not None:
        rows = [me]
    else:
        rows = []

    return ApiResponse(
        data={"companies": [{"id": c.id, "name": c.name} for c in rows]},
        message="",
    )


# ═══════════════════════════════════════════════════════════════
# 人员库（本公司）
# ═══════════════════════════════════════════════════════════════

@router.get("/personnel", response_model=ApiResponse)
async def list_personnel(
    with_cert_only: bool = Query(False, description="只看持稳评培训证书的"),
    roster_only: bool = Query(False, description="只列公司名册上的在职人员（默认**关**）"),
    names: str = Query("", description="逗号分隔的人名，只取这些人（用于「已选人员」回显）"),
    user: User = Depends(get_current_user),
):
    """**本公司**人员库 —— 供「选择报告人员」多选框使用。

    🔴 `roster_only`（默认**关**）：名册（`company_assets`「公司专业技术人员配备」）
       是投标文件的旧快照，只有 8 人、且只有 3 人那栏填了稳评证书；而库里**实际有
       17 人持稳评培训证书**（证书扫描件为证）。按名册过滤会把 14 个持证的人藏起来
       （业主：「持证人员可选择的太少」）。要看名册就传 `roster_only=true`。
    """
    from app.services.personnel_store import images_for, personnel_for, roster_names

    cid = user.company_id
    want = [n.strip() for n in names.split(",") if n.strip()] or None
    people = personnel_for(want, cid)
    _roster = roster_names(cid) if cid else set()

    img_cache: dict = {}
    out = []
    for p in people:
        if with_cert_only and not p.get("has_wenshen_cert"):
            continue
        name = p["姓名"]
        if roster_only and _roster and name not in _roster:
            continue
        if name not in img_cache:
            img_cache[name] = [cap for _, cap in images_for([name], cid)]
        out.append({
            # 🔴 必须把人员 id 给前端 —— 修改/删除/上传证书都要用它定位
            "id": p.get("_id"),
            "name": name,
            "in_roster": name in _roster,
            "gender": p.get("性别", ""),
            "age": p.get("年龄", ""),
            "education": p.get("学历", ""),
            "major": p.get("专业", ""),
            "title": p.get("职称/职务", ""),
            "certs": p.get("执业资格证书", []),
            "default_role": p.get("拟任角色", "项目组成员"),
            "has_wenshen_cert": bool(p.get("has_wenshen_cert")),
            "cert_images": img_cache[name],
            "cert_image_count": len(img_cache[name]),
        })

    # 按「持证 → 负责人 → 其他」排，方便前端默认勾选
    out.sort(key=lambda x: (not x["has_wenshen_cert"], x["default_role"] != "项目负责人", x["name"]))
    msg = f"共 {len(out)} 名人员" if cid else "尚未填写公司信息，人员库为空"
    return ApiResponse(message=msg, data={"personnel": out, "total": len(out),
                                          "has_company": bool(cid)})


@router.post("/personnel/sync", response_model=ApiResponse)
async def sync_personnel(
    images: bool = Form(True, description="是否同时重建人员图片绑定"),
    user: User = Depends(get_current_user),
):
    """从**本公司名册 + 素材图**重建人员库（幂等，可反复执行）。

    ⚠️ 名册来源（`company_assets`）目前只有内部公司那一份 —— 外部用户
       应当走 `POST /personnel/import` 上传自己的名册。
    """
    from app.services.personnel_store import sync_personnel_store
    if not user.company_id:
        raise HTTPException(status_code=400, detail="请先填写公司信息")
    n = sync_personnel_store(rebuild_images=images, company_id=user.company_id)
    return ApiResponse(message=f"✅ 人员库已更新：{n['personnel']} 人 / {n['images']} 张图片", data=n)


# 名册表头 → 人员库字段（`personnel_store` 的中文键，必须严格对齐，否则下游取不到）
_ROSTER_COLUMN_MAP = {
    "姓名": "姓名", "名字": "姓名",
    "性别": "性别", "年龄": "年龄", "学历": "学历", "专业": "专业",
    "职称": "职称", "职务": "职称", "职称/职务": "职称",
    "拟任角色": "拟任角色", "角色": "拟任角色",
    "执业资格证书": "执业资格证书", "证书": "执业资格证书",
}


def _parse_roster_rows(raw_bytes: bytes, filename: str) -> List[dict]:
    """把上传的名册（xlsx/csv/docx）读成 [{中文键: 值}]。

    只认表头里出现过的列，认不出的列忽略 —— 用户的名册格式千奇百怪，
    硬要求某一种列名只会让人传不上来。
    """
    name = (filename or "").lower()

    def _rows_from_table(rows):
        if not rows:
            return []
        header = [str(c or "").strip() for c in rows[0]]
        idx = {}
        for i, h in enumerate(header):
            key = _ROSTER_COLUMN_MAP.get(h.replace(" ", ""))
            if key and key not in idx:
                idx[key] = i
        if "姓名" not in idx:
            return []
        out = []
        for r in rows[1:]:
            item = {}
            for key, i in idx.items():
                v = r[i] if i < len(r) else ""
                item[key] = "" if v is None else str(v).strip()
            if item.get("姓名"):
                out.append(item)
        return out

    if name.endswith(".csv") or name.endswith(".txt"):
        import csv
        import io
        text = raw_bytes.decode("utf-8-sig", errors="ignore")
        return _rows_from_table([row for row in csv.reader(io.StringIO(text))])

    if name.endswith(".xlsx") or name.endswith(".xls"):
        try:
            from openpyxl import load_workbook
            import io
            wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            return _rows_from_table([list(r) for r in ws.iter_rows(values_only=True)])
        except Exception as e:
            logger.warning(f"名册 xlsx 解析失败: {e}")
            return []

    if name.endswith(".docx"):
        try:
            from docx import Document
            import io
            d = Document(io.BytesIO(raw_bytes))
            for t in d.tables:                      # 名册通常就是一张表
                rows = _rows_from_table([[c.text for c in r.cells] for r in t.rows])
                if rows:
                    return rows
            return []
        except Exception as e:
            logger.warning(f"名册 docx 解析失败: {e}")
            return []

    return []


@router.post("/personnel/import", response_model=ApiResponse)
async def import_personnel(
    file: UploadFile = File(..., description="名册：xlsx / csv / docx（第一行是表头）"),
    replace: bool = Form(False, description="先清空本公司现有人员再导入"),
    user: User = Depends(get_current_user),
):
    """上传**自己的**人员名册 → 写入本公司人员库。

    表头认「姓名/性别/年龄/学历/专业/职称/拟任角色/执业资格证书」这几个键
    （与 `personnel_store` 的字段严格对齐），认不出的列忽略。
    """
    from app.services.personnel_store import _conn, ensure_tables

    if not user.company_id:
        raise HTTPException(status_code=400, detail="请先填写公司信息")

    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="文件为空")
    rows = _parse_roster_rows(raw, file.filename or "")
    if not rows:
        raise HTTPException(
            status_code=400,
            detail="没解析出人员。请确认第一行是表头且含「姓名」列（支持 xlsx / csv / docx）",
        )

    ensure_tables()
    c = _conn()
    if not c:
        raise HTTPException(status_code=500, detail="人员库不可用")
    try:
        if replace:
            c.execute("DELETE FROM company_personnel WHERE company_id = ?", (user.company_id,))
        added = 0
        for i, r in enumerate(rows):
            certs = [x for x in str(r.get("执业资格证书") or "").replace("；", ",")
                     .replace("、", ",").split(",") if x.strip()]
            c.execute(
                "INSERT INTO company_personnel "
                "(name, gender, age, education, major, title, role, certs, "
                " has_wenshen_cert, sort_order, company_id, updated_at) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?, datetime('now')) "
                "ON CONFLICT(name, company_id) DO UPDATE SET "
                "gender=excluded.gender, age=excluded.age, education=excluded.education, "
                "major=excluded.major, title=excluded.title, role=excluded.role, "
                "certs=excluded.certs, has_wenshen_cert=excluded.has_wenshen_cert, "
                "updated_at=excluded.updated_at",
                (r["姓名"], r.get("性别", ""), r.get("年龄", ""), r.get("学历", ""),
                 r.get("专业", ""), r.get("职称", ""),
                 r.get("拟任角色") or "项目组成员", json.dumps(certs, ensure_ascii=False),
                 1 if any("培训" in x for x in certs) else 0, i, user.company_id),
            )
            added += 1
        c.commit()
        return ApiResponse(message=f"✅ 已导入 {added} 名人员", data={"imported": added})
    except Exception as e:
        c.rollback()
        logger.exception("名册导入失败")
        raise HTTPException(status_code=500, detail=f"导入失败：{e}")
    finally:
        c.close()


@router.post("/personnel", response_model=ApiResponse)
async def create_personnel(
    payload: dict,
    user: User = Depends(get_current_user),
):
    """手工新增**本公司**一名人员（名册导入之外的补充入口）。"""
    from app.services.personnel_store import _conn, ensure_tables

    if not user.company_id:
        raise HTTPException(status_code=400, detail="请先填写公司信息")
    name = str(payload.get("name") or payload.get("姓名") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="请填写姓名")

    ensure_tables()
    c = _conn()
    if not c:
        raise HTTPException(status_code=500, detail="人员库不可用")
    try:
        certs = payload.get("certs") or payload.get("执业资格证书") or []
        if isinstance(certs, str):
            certs = [x for x in certs.replace("；", ",").replace("、", ",").split(",") if x.strip()]
        c.execute(
            "INSERT INTO company_personnel "
            "(name, gender, age, education, major, title, role, certs, has_wenshen_cert, "
            " sort_order, company_id, updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?, datetime('now')) "
            "ON CONFLICT(name, company_id) DO UPDATE SET "
            "gender=excluded.gender, age=excluded.age, education=excluded.education, "
            "major=excluded.major, title=excluded.title, role=excluded.role, "
            "certs=excluded.certs, has_wenshen_cert=excluded.has_wenshen_cert, "
            "updated_at=excluded.updated_at",
            (name, str(payload.get("gender") or ""), str(payload.get("age") or ""),
             str(payload.get("education") or ""), str(payload.get("major") or ""),
             str(payload.get("title") or ""), str(payload.get("role") or "项目组成员"),
             json.dumps([str(x) for x in certs], ensure_ascii=False),
             1 if any("培训" in str(x) for x in certs) else 0,
             999, user.company_id),
        )
        c.commit()
        return ApiResponse(message=f"✅ 已添加 {name}")
    except Exception as e:
        c.rollback()
        logger.exception("新增人员失败")
        raise HTTPException(status_code=500, detail=f"新增失败：{e}")
    finally:
        c.close()


@router.put("/personnel/{personnel_id}", response_model=ApiResponse)
async def update_personnel(
    personnel_id: int,
    payload: dict,
    user: User = Depends(get_current_user),
):
    """修改本公司某位人员的信息（只改传进来的字段）。"""
    from app.services.personnel_store import _conn, ensure_tables

    # 🔴 这个判据只对普通账号有意义：管理员可以**跨公司**操作（运维代客户处理），
    #    他自己没有公司也照样能改。
    if not user.company_id and not is_admin(user):
        raise HTTPException(status_code=400, detail="请先填写公司信息")
    ensure_tables()
    c = _conn()
    if not c:
        raise HTTPException(status_code=500, detail="人员库不可用")
    try:
        # 🔴 管理员可跨公司改（运维/客服代客户处理）；普通账号只能改本公司的人
        if is_admin(user):
            row = c.execute("SELECT * FROM company_personnel WHERE id = ?",
                            (personnel_id,)).fetchone()
            _scope_cid = row["company_id"] if row else None
        else:
            row = c.execute("SELECT * FROM company_personnel WHERE id = ? AND company_id = ?",
                            (personnel_id, user.company_id)).fetchone()
            _scope_cid = user.company_id
        if not row:
            raise HTTPException(status_code=404, detail="人员不存在")

        field_map = {
            "name": "name", "姓名": "name",
            "gender": "gender", "性别": "gender",
            "age": "age", "年龄": "age",
            "education": "education", "学历": "education",
            "major": "major", "专业": "major",
            "title": "title", "职称/职务": "title", "title_name": "title",
            "role": "role", "拟任角色": "role",
        }
        sets, params = [], []
        for k, col in field_map.items():
            if k in payload and payload[k] is not None:
                sets.append(f"{col} = ?"); params.append(str(payload[k]).strip())
        if "certs" in payload or "执业资格证书" in payload:
            certs = payload.get("certs", payload.get("执业资格证书")) or []
            if isinstance(certs, str):
                certs = [x for x in certs.replace("；", ",").replace("、", ",").split(",") if x.strip()]
            certs = [str(x) for x in certs]
            sets += ["certs = ?", "has_wenshen_cert = ?"]
            params += [json.dumps(certs, ensure_ascii=False),
                       1 if any("培训" in x for x in certs) else 0]

        if not sets:
            return ApiResponse(message="没有需要修改的字段")
        sets.append("updated_at = datetime('now')")
        params += [personnel_id, _scope_cid]
        c.execute(f"UPDATE company_personnel SET {', '.join(sets)} "
                  f"WHERE id = ? AND company_id = ?", params)
        c.commit()
        return ApiResponse(message="✅ 已保存")
    except HTTPException:
        raise
    except Exception as e:
        c.rollback()
        # 同名冲突（公司内唯一）走这里，给个能看懂的提示
        if "UNIQUE" in str(e).upper():
            raise HTTPException(status_code=409, detail="本公司已有同名人员")
        logger.exception("修改人员失败")
        raise HTTPException(status_code=500, detail=f"保存失败：{e}")
    finally:
        c.close()


@router.delete("/personnel/{personnel_id}", response_model=ApiResponse)
async def delete_personnel(
    personnel_id: int,
    user: User = Depends(get_current_user),
):
    """删除本公司某位人员（连同其证书图片记录）。

    ⚠️ 只删库里的记录，磁盘上的证书图不删 —— 避免误删还可能被别处引用的文件。
    """
    from app.services.personnel_store import _conn, ensure_tables

    if not user.company_id and not is_admin(user):
        raise HTTPException(status_code=400, detail="请先填写公司信息")
    ensure_tables()
    c = _conn()
    if not c:
        raise HTTPException(status_code=500, detail="人员库不可用")
    try:
        # 🔴 同上：管理员可跨公司
        if is_admin(user):
            row = c.execute("SELECT name, company_id FROM company_personnel WHERE id = ?",
                            (personnel_id,)).fetchone()
        else:
            row = c.execute("SELECT name, company_id FROM company_personnel WHERE id = ? AND company_id = ?",
                            (personnel_id, user.company_id)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="人员不存在")
        _cid = row["company_id"]
        c.execute("DELETE FROM company_personnel_images WHERE personnel_id = ? AND company_id = ?",
                  (personnel_id, _cid))
        c.execute("DELETE FROM company_personnel WHERE id = ? AND company_id = ?",
                  (personnel_id, _cid))
        c.commit()
        return ApiResponse(message=f"已删除 {row['name']}")
    except HTTPException:
        raise
    except Exception as e:
        c.rollback()
        logger.exception("删除人员失败")
        raise HTTPException(status_code=500, detail=f"删除失败：{e}")
    finally:
        c.close()


@router.post("/personnel/{personnel_id}/certs", response_model=ApiResponse)
async def upload_personnel_cert(
    personnel_id: int,
    file: UploadFile = File(..., description="证书图片（png/jpg）"),
    category: str = Form("其他", description="培训证书 / 职称证书 / 执业资格证书 / 学历证书 / 证件照 / 其他"),
    user: User = Depends(get_current_user),
):
    """给本公司某位人员上传一张证书图片。"""
    from app.services.personnel_store import _conn, ensure_tables
    from app.config import settings

    if not user.company_id:
        raise HTTPException(status_code=400, detail="请先填写公司信息")

    ensure_tables()
    c = _conn()
    if not c:
        raise HTTPException(status_code=500, detail="人员库不可用")
    try:
        # 只能给本公司的挂证书
        row = c.execute(
            "SELECT id, name FROM company_personnel WHERE id = ? AND company_id = ?",
            (personnel_id, user.company_id)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="人员不存在")

        ext = (file.filename or "").rsplit(".", 1)[-1].lower()
        if ext not in ("png", "jpg", "jpeg", "webp", "bmp"):
            raise HTTPException(status_code=400, detail="只支持图片（png/jpg/webp/bmp）")

        import uuid
        dest_dir = settings.STORAGE_DIR / "companies" / str(user.company_id)
        dest_dir.mkdir(parents=True, exist_ok=True)
        fname = f"{row['name']}_{uuid.uuid4().hex[:8]}.{ext}"
        dest = dest_dir / fname
        dest.write_bytes(await file.read())

        c.execute(
            "INSERT INTO company_personnel_images "
            "(personnel_id, image_name, category, image_path, mime_type, source, company_id, updated_at) "
            "VALUES (?,?,?,?,?,?,?, datetime('now')) "
            "ON CONFLICT(personnel_id, image_name) DO UPDATE SET "
            "category=excluded.category, image_path=excluded.image_path, updated_at=excluded.updated_at",
            (personnel_id, fname, category, str(dest),
             f"image/{ext}", "user_upload", user.company_id),
        )
        c.commit()
        return ApiResponse(message=f"✅ 已上传 {row['name']} 的{category}", data={"path": str(dest)})
    except HTTPException:
        raise
    except Exception as e:
        c.rollback()
        logger.exception("证书上传失败")
        raise HTTPException(status_code=500, detail=f"上传失败：{e}")
    finally:
        c.close()

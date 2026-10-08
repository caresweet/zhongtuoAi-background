"""回归测试：评审表模板选取 + 参照件照搬章渲染。

锁住两个已修的真实缺陷，防止回退：

① **评审表模板选错**（曾把整份投标文件搬进报告第三章、400+ 页）
   旧实现裸 `sorted(glob("*.docx"))` + `any(tbl)` 判据 → 字典序第一的
   `18f2827d…docx`（36MB 投标文件，区间 1125 元素/9 张表）被选中。
   现在必须：洪泽模板优先、且区间**恰为 1 张 22×7 制式表**。

② **参照件照搬章**：必须零 LLM、无占位符残留、无参照件专有词泄漏。
"""

import os
import re
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

_TEMPLATES = BACKEND / "storage" / "templates"
_BID_DOC = "18f2827d62fe40e48cc71b58547db187.docx"   # 投标文件（绝不能选中）


# ════════════════════════════════════════════════════════
# ① 评审表模板选取
# ════════════════════════════════════════════════════════

@pytest.mark.skipif(not _TEMPLATES.exists(), reason="无模板目录")
def test_bid_document_is_rejected_as_review_table():
    """投标文件（1125 元素/9 张表）必须被判为非评审表区间。"""
    p = _TEMPLATES / _BID_DOC
    if not p.exists():
        pytest.skip("投标文件样本不在盘上")
    from app.services.review_table_filler import _template_review_region, _is_review_region
    _tpl, region = _template_review_region(str(p))
    assert len(region) > 100, "投标文件区间本就很大（用于验证判据有效性）"
    assert not _is_review_region(region), "投标文件不得被判为评审表区间"


@pytest.mark.skipif(not _TEMPLATES.exists(), reason="无模板目录")
def test_rank_prefers_stability_template():
    """候选排序必须把洪泽稳评模板排在任何投标/磋商文件之前。"""
    from app.services.template_chapter_chunker import rank_template_candidates
    cands = rank_template_candidates()
    if not cands:
        pytest.skip("无候选模板")
    assert "598da27a" in os.path.basename(cands[0]), \
        f"首个候选应为洪泽稳评模板，实际是 {os.path.basename(cands[0])}"
    if any(_BID_DOC in c for c in cands):
        i_bid = next(i for i, c in enumerate(cands) if _BID_DOC in c)
        assert i_bid > 0, "投标文件不得排在首位"


@pytest.mark.skipif(not _TEMPLATES.exists(), reason="无模板目录")
def test_region_validator_accepts_only_22x7():
    """结构判据：只有「恰 1 张 22 行 × 7 栅格」的区间才通过。"""
    from app.services.review_table_filler import _template_review_region, _is_review_region
    accepted = []
    for p in _TEMPLATES.glob("*.docx"):
        try:
            _tpl, region = _template_review_region(str(p))
        except Exception:
            continue
        if _is_review_region(region):
            tbls = [c for c in region if c.tag.split('}')[1] == 'tbl']
            assert len(tbls) == 1
            accepted.append(p.name)
    # 至少洪泽模板要通过；且投标文件必须不在内
    assert any("598da27a" in n for n in accepted), f"洪泽模板未被接受：{accepted}"
    assert _BID_DOC not in accepted


# ════════════════════════════════════════════════════════
# ② 参照件照搬章
# ════════════════════════════════════════════════════════

_FILLED = {
    "project_name": "洪拟征告〔2026〕7号商业服务业设施用地项目",
    "org_name": "淮安市洪泽区人民政府",
    "decision_unit": "淮安市洪泽区人民政府",
    "location": "朱坝街道办事处、三圩社区、二组、三组、六组",
    "area_mu": "489.51",
    "area_m2": "326342",
    "land_use": "商业服务业设施用地",
    "doc_reference": "洪拟征告〔2026〕7号",
}


def _corpus_available() -> bool:
    from app.services.reference_corpus import is_enabled
    try:
        return is_enabled()
    except Exception:
        return False


@pytest.mark.skipif(not _corpus_available(), reason="参照件语料未生成")
def test_verbatim_chapters_render_clean():
    """第1/2/5/7章渲染后：无占位符残留、无参照件专有词泄漏。"""
    from app.services.reference_corpus import render_reference_markdown
    for n in (1, 2, 5, 7):
        r = render_reference_markdown(n, _FILLED)
        assert r.markdown.strip(), f"第{n}章渲染为空"
        assert not re.findall(r'\{\{\w+\}\}', r.markdown), f"第{n}章有未替换占位符"
        assert "【待补充" not in r.markdown, f"第{n}章残留占位符标记"
        assert not r.leaks, f"第{n}章泄漏参照件专有词：{r.leaks}"


@pytest.mark.skipif(not _corpus_available(), reason="参照件语料未生成")
def test_verbatim_agent_never_calls_llm(monkeypatch):
    """照搬模式（REFERENCE_LLM_REWRITE=False）必须零 LLM 调用。"""
    import asyncio
    from app.config import settings
    from app.agent.agents.chapters import get_chapter_agent
    from app.services.reference_corpus import verbatim_chapters

    monkeypatch.setattr(settings, "REFERENCE_LLM_REWRITE", False, raising=False)

    class _Boom:
        def __getattr__(self, k):
            raise AssertionError(f"照搬模式不应调用 LLM（访问了 {k}）")

    for n in sorted(verbatim_chapters()):
        agent = get_chapter_agent(n, llm_service=_Boom())
        state = {"filled_data": dict(_FILLED), "chapters": {}}
        q = asyncio.Queue()
        asyncio.run(agent.run(state, q))
        ch = state["chapters"][n]
        assert ch.get("status") == "reference_verbatim"
        assert ch["markdown"].strip()


@pytest.mark.skipif(not _corpus_available(), reason="参照件语料未生成")
def test_guided_agent_rewrites_with_llm(monkeypatch):
    """改写模式（REFERENCE_LLM_REWRITE=True）：正文由模型写，且骨架结构一个不丢。

    业主口径 2026-09-16：「除项目事实外全交给模型」。
    """
    import asyncio
    from app.config import settings
    from app.agent.agents.chapters import get_chapter_agent

    monkeypatch.setattr(settings, "REFERENCE_LLM_REWRITE", True, raising=False)

    calls = {"n": 0, "prompt": "", "system": ""}

    class _FakeLLM:
        async def chat_with_reasoning(self, messages=None, system=None, **kw):
            calls["n"] += 1
            calls["prompt"] = (messages or [{}])[0].get("content", "")
            calls["system"] = system or ""
            # 取提示词里骨架围栏内的标题/图/表标记，每节下面写本项目自己的新正文
            fence = re.search(r'```markdown\n(.*?)\n```', calls["prompt"], flags=re.S)
            anchors = [l for l in (fence.group(1) if fence else "").split("\n") if l.strip()]
            filler = ("本项目严格按照上述要求组织实施，明确责任分工与办理时限，"
                      "确保各项措施落实到位、风险可控。") * 8
            return {"content": "\n\n".join(
                f"{a}\n\n{filler}" if a.startswith("#") else a for a in anchors)}

    agent = get_chapter_agent(5, llm_service=_FakeLLM())
    state = {"filled_data": dict(_FILLED), "chapters": {}, "session_id": ""}
    asyncio.run(agent.run(state, asyncio.Queue()))
    ch = state["chapters"][5]

    assert calls["n"] == 1, "改写模式必须调用模型一次"
    assert ch.get("status") == "reference_guided"
    assert ch.get("llm_rewritten") is True
    # 骨架结构必须原样保留（标题 / 插图 / 系统表标记）
    from app.services.reference_llm import _anchors
    from app.services.reference_corpus import render_reference_markdown
    assert _anchors(ch["markdown"]) == _anchors(render_reference_markdown(5, dict(_FILLED)).markdown)
    # 提示词里必须带上禁令与素材
    assert "绝对不许编造" in calls["system"], "系统提示必须带禁令"
    assert "项目事实卡" in calls["prompt"] or "事实卡" in calls["prompt"], "用户提示必须带事实卡"
    assert "可引用法规清单" in calls["prompt"], "用户提示必须带法规白名单"


@pytest.mark.skipif(not _corpus_available(), reason="参照件语料未生成")
def test_guided_rewrite_falls_back_when_structure_broken(monkeypatch):
    """模型把标题/图吞了 → 必须回退照搬稿，绝不放出结构损坏的章节。"""
    import asyncio
    from app.config import settings
    from app.agent.agents.chapters import get_chapter_agent

    monkeypatch.setattr(settings, "REFERENCE_LLM_REWRITE", True, raising=False)

    class _Mangler:
        async def chat_with_reasoning(self, messages=None, system=None, **kw):
            return {"content": "### （一）总则\n\n" + "随便写了一段很长的正文。" * 30}

    agent = get_chapter_agent(5, llm_service=_Mangler())
    state = {"filled_data": dict(_FILLED), "chapters": {}, "session_id": ""}
    asyncio.run(agent.run(state, asyncio.Queue()))
    ch = state["chapters"][5]
    assert ch.get("status") == "reference_verbatim", "结构损坏时必须回退照搬稿"


# ── 第二层校验：不许编的东西 ──────────────────────────────

def test_guard_strips_phone_and_unlisted_docref():
    """手机号与非事实卡文号必须硬剥；白名单外法规与大额数字必须报出来。"""
    from app.services.reference_llm import guard_fabrications

    md = ("联系人为张三，电话13812345678。\n"
          "依据苏办发〔2021〕15号开展评估。\n"
          "另参照洪拟征告〔2026〕7号执行。\n"
          "依据《中华人民共和国突发事件应对法》办理。\n"
          "依据《虚构的风险评估条例》办理。\n"
          "项目投资约 45000 万元。")
    out, issues = guard_fabrications(md, dict(_FILLED), skeleton="")
    kinds = [i["kind"] for i in issues]

    assert "13812345678" not in out, "手机号必须剥除"
    assert "苏办发〔2021〕15号" not in out, "非事实卡文号必须剥除"
    assert "洪拟征告〔2026〕7号" in out, "事实卡里的文号必须保留"
    assert "phone_stripped" in kinds
    assert "docref_stripped" in kinds
    assert "regulation_unlisted" in kinds, "白名单外法规必须报出"
    assert any("45000" in i.get("detail", "") for i in issues), "未见于事实卡的大额数字必须报出"


def test_structure_ok_detects_missing_and_extra_anchors():
    from app.services.reference_llm import structure_ok

    sk = "### （一）总则\n\n正文\n\n![流程图](figures/fig5-1.jpeg)\n\n#### 1.编制目的\n\n正文"
    assert structure_ok(sk, sk)[0]
    ok, why = structure_ok("### （一）总则\n\n正文", sk)
    assert not ok and "缺少" in why
    ok, why = structure_ok(sk + "\n\n### （十三）多出来的", sk)
    assert not ok and "多出" in why


def test_law_paragraphs_are_protected_verbatim():
    """第一档 · 法规条文段：进骨架、必须逐字保留，被「顺手润色」要判失败。

    业主口径（2026-09-17）：地区政策法规**可以百分百 copy**，不需要改。
    """
    from app.services.reference_llm import (
        _is_law_paragraph, _protected, _skeleton_only, structure_ok,
    )

    law1 = "依据《中华人民共和国突发事件应对法》《信访工作条例》等法律法规和有关规定，制定本预案。"
    law2 = "1.《中华人民共和国土地管理法》；"
    prose = "本项目坚持预防为主，把矛盾化解在基层，切实维护群众合法权益和社会稳定大局。"

    assert _is_law_paragraph(law1) and _is_law_paragraph(law2)
    assert not _is_law_paragraph(prose), "普通论述段不算条文，那是第三档要仿照重写的"

    skeleton = f"### （二）编制依据\n\n{law1}\n\n{prose}\n\n{law2}"
    skel = _skeleton_only(skeleton)
    assert law1 in skel and law2 in skel, "法规条文段必须进骨架"
    assert "把矛盾化解在基层" not in skel, "普通论述段不该进骨架（留给模型自己写）"
    assert len(_protected(skeleton)) == 2

    assert structure_ok(skeleton, skeleton)[0]
    ok, why = structure_ok(skeleton.replace("制定本预案", "制定本方案"), skeleton)
    assert not ok and "法规条文段" in why


@pytest.mark.skipif(not _corpus_available(), reason="参照件语料未生成")
def test_verbatim_excludes_wind_content():
    """风电专属法规/条目必须被剔除（参照件是风电项目，本项目是土地征收）。"""
    from app.services.reference_corpus import render_reference_markdown
    md = render_reference_markdown(2, _FILLED).markdown
    for kw in ("陆上风电", "驭风", "可再生能源发展专项规划", "千乡万村"):
        assert kw not in md, f"第2章仍含风电专属条目「{kw}」"


@pytest.mark.skipif(not _corpus_available(), reason="参照件语料未生成")
def test_verbatim_chapter5_keeps_two_level_hierarchy():
    """第5章的二级/三级结构必须与参照件一致（业主：「依据模板来」）。

    参照件 5.1 总则 下辖 5.1.1 编制目的 / 5.1.2 编制依据 / 5.1.3 适用范围。
    旧实现把所有 `#{2,4} 数字…` 一律压成同级的 `### （N）`，
    于是 15 条标题被拉平成 （一）…（十五），模板层级整段丢失。
    """
    from app.services.reference_corpus import render_reference_markdown
    md = render_reference_markdown(5, _FILLED).markdown
    h2 = [l.strip() for l in md.split("\n") if l.strip().startswith("### ")]
    h3 = [l.strip() for l in md.split("\n") if l.strip().startswith("#### ")]
    assert len(h2) == 12, f"二级标题应为 12 条，实得 {len(h2)}: {h2}"
    # 三级挂在正确的二级下面
    assert h3[:3] == ["#### 1.编制目的", "#### 2.编制依据", "#### 3.适用范围"]
    assert "#### 1.预警信息" in h3
    assert "#### 1.通信保障" in h3
    # 粘连标题必须被拆开，不能整条留在二级里
    assert not any("1.预警信息" in x for x in h2), "「预警机制1.预警信息」未拆开"
    assert not any("1.通信保障" in x for x in h2), "「保障措施1.通信保障」未拆开"
    # 三级序号在各自的二级内重排
    assert "### （一）总则" in md and "### （十二）奖惩、责任追究" in md


@pytest.mark.skipif(not _corpus_available(), reason="参照件语料未生成")
def test_verbatim_chapters_have_no_personal_data():
    """参照件里他人姓名/手机号绝不进本项目成稿（隐私）。"""
    from app.services.reference_corpus import render_reference_markdown
    for n in (1, 2, 5, 7):
        md = render_reference_markdown(n, _FILLED).markdown
        assert not re.findall(r'(?<!\d)1[3-9]\d{9}(?!\d)', md), f"第{n}章残留手机号"
    # 第5章「组织领导」的名单段必须已去个人化
    md5 = render_reference_markdown(5, _FILLED).markdown
    assert "组长：" not in md5, "第5章仍残留他人领导小组名单"


def test_image_caption_survives_post_generation_fix(tmp_path):
    """图注不能在后处理里被清空，且图号必须按**本章**章号编（业主：「流程图没有」）。

    两个真实缺陷：
    ① `_post_generation_fix` 用 `id(Paragraph)` 认「图注段落」，而 `doc.paragraphs`
       每次都新建代理对象、地址会被复用 → 图注被当普通段落清空（整条图注消失）；
    ② 章号由 `assemble()` 循环设置，单章导出直接调 `_add_chapter` → 全编成「图0-N」。
    """
    from docx import Document
    from PIL import Image
    from app.services.report_assembler import ReportAssembler

    img = tmp_path / "flow.png"
    Image.new("RGB", (400, 300), "white").save(img)

    a = ReportAssembler()
    a._inserted_images = set(); a._chapter_image_counter = {}
    a._system_rendered_tables = set(); a._rendered_extracted_tbls = set()
    a._session_state = {}
    d = Document(); a._setup_page(d)
    a._add_chapter(d, 5, f"### （六）工作流程\n\n如下图所示：\n\n![社会稳定风险处理流程图]({img})",
                   {}, {}, [], {})

    _want = ["图5-1 社会稳定风险处理流程图"]
    assert [p.text.strip() for p in d.paragraphs if "流程图" in p.text] == _want
    a._post_generation_fix(d, {})
    assert [p.text.strip() for p in d.paragraphs if "流程图" in p.text] == _want, "图注被后处理清空"


# ════════════════════════════════════════════════════════
# ③ RAG 必须真正进入提示词（业主：「不要让 LLM 乱编，要结合 RAG」）
# ════════════════════════════════════════════════════════

def test_custom_prompt_path_still_injects_rag():
    """走自定义 prompt 的路径（脚本生成全走这条）也必须带上 RAG 素材。

    此前该分支直接 `user_prompt = state["_custom_prompt"]`，
    把 think() 检索到的知识库素材**整个丢掉**，LLM 只能凭记忆编法规名与规划名。
    """
    from app.agent.agents.chapters.chapter_4_agent import Chapter4Agent
    agent = Chapter4Agent()
    rag = {
        "chapter_context": "《中华人民共和国土地管理法》第四十五条规定……" * 3,
        "four_dimension_context": "【8.1.2 地方规划相符性】符合《XX区国土空间分区规划》……" * 3,
        "sources": [{"x": 1}],
    }
    block = agent._format_rag_for_prompt(rag)
    assert block, "RAG 素材应当被渲染进提示词"
    assert "四性分析" in block, "四性分析专线素材必须出现在素材块里"
    assert "只能" in block and "不许" in block, "必须带「只能用素材、不得编造」的硬约束"
    assert "范例" in block, "必须说明范例只作参考、不得抄其项目名/数字"


def test_four_dimension_chapter_uses_dedicated_retrieval():
    """第4章/第10章必须走四性分析专线（主类型取回的是评分表碎片，取不到四性范文）。"""
    from app.agent.agents.chapters.chapter_base import ChapterAgentBase
    assert hasattr(ChapterAgentBase, "_retrieve_four_dimension")
    import inspect
    src = inspect.getsource(ChapterAgentBase._retrieve_rag)
    assert "four_dimension_context" in src, "四性分析专线未接进 _retrieve_rag"


def test_section_subs_normalized_to_headings():
    """二级标题体例统一为（一）（二）（三），且分条表不得再带 X.Y.Z 前缀。

    🔴 业主口径「全报告体例照参照件 1号」：二级标题一律（一）（二）（三），
    块内再分条用 1. 2. 3.。此前是 1.1/1.2 与 4.2.1 三级编号。
    """
    import re as _re
    from app.agent.chapter_definitions import (
        CHAPTER_DEFINITIONS, get_chapter_section_subs)
    for n, d in CHAPTER_DEFINITIONS.items():
        for sec in (d.get("sections") or []):
            assert _re.match(r'^（[一二三四五六七八九十]+）', sec), f"第{n}章二级标题体例不对：{sec}"
    for n in CHAPTER_DEFINITIONS:
        subs = get_chapter_section_subs(n)
        for _sec, items in (subs or {}).items():
            for it in items:
                assert not _re.match(r'^\d+(?:\.\d+)+', it), f"分条仍带 X.Y.Z 前缀：{it}"


# ════════════════════════════════════════════════════════
# ⑦ 第1章：三块照搬原文，只扩写前言
# ════════════════════════════════════════════════════════

def test_chapter1_structure_is_locked():
    """第1章结构锁定为「（一）稳评理由 /（二）稳评目的 /（三）稳评原则」三块。"""
    from app.services import chapter1_composer as C1
    _pre, blocks = C1.build_blocks(dict(_FILLED))
    assert [b for b, _ in blocks] == ["（一）稳评理由", "（二）稳评目的", "（三）稳评原则"]
    items = {b: [lab for lab, _ in its] for b, its in blocks}
    assert len(items["（一）稳评理由"]) == 3
    assert len(items["（二）稳评目的"]) == 3
    assert len(items["（三）稳评原则"]) == 4
    for labs in items.values():
        for lab in labs:
            assert lab[0].isdigit() and lab[1] in ".、", f"条目编号异常：{lab}"


def test_chapter1_blocks_are_verbatim_not_rewritten():
    """🔴 三块内容必须**原文照搬**，只有前言可以补写。

    业主口径：「我让你添加内容的模块是根据中共中央那块的内容，
    不是（一）稳评理由/（二）稳评目的/（三）稳评原则这三块。」
    """
    from app.services import chapter1_composer as C1
    pre, blocks = C1.build_blocks(dict(_FILLED))
    md = C1.compose(pre, blocks, ["这是前言补写的第一段。", "这是前言补写的第二段。"])
    # 三块与条目名必须在
    for b, its in blocks:
        assert f"### {b}" in md, f"缺少结构块 {b}"
        for lab, body in its:
            assert lab in md, f"条目名丢失：{lab}"
            if body:
                assert body[:30] in md, f"条目的原说明段被改写：{lab}"
    # 补写段落只出现在第一个块之前（即前言区）
    i_pre = md.find("### （一）稳评理由")
    for extra in ("这是前言补写的第一段。", "这是前言补写的第二段。"):
        assert md.find(extra) < i_pre, f"补写段落跑到了三块内部：{extra}"


def test_chapter1_compose_has_no_extra_inside_blocks():
    """不传补写段时，产物就是纯粹的「前言 + 三块原文」。"""
    from app.services import chapter1_composer as C1
    pre, blocks = C1.build_blocks(dict(_FILLED))
    md = C1.compose(pre, blocks, [])
    assert "### （一）稳评理由" in md and "### （三）稳评原则" in md
    assert md.count("### （一）稳评理由") == 1


# ════════════════════════════════════════════════════════
# ⑧ 组装器的两个静默丢内容缺陷
# ════════════════════════════════════════════════════════

def test_short_lines_survive_length_guard():
    """长度守卫曾静默吞内容 —— 现在判据是「有没有实义」，不是「够不够长」。

    实测被吞掉的两类：
      ①「1.客观公正性原则」（9 字）—— 正文只剩四条说明段，原则名全丢；
      ② 第7章落款「稳评责任单位（盖章）」（正好 10 字）与「年月日」——
        「3.方案审核」后面直接接第八章，落款整块消失。
    倾向是**宁可多留**：多留一行人工一眼能看出，少一行没人看得出来。
    """
    from app.services.report_assembler import _has_meaning as ok
    for t in ("1.客观公正性原则", "2.科学评价性原则", "（一）稳评理由",
              "A. 要点", "图5-1", "表12-3", "第一章 编制说明",
              "稳评责任单位（盖章）", "年月日", "2026年9月13日"):
        assert ok(t), f"有实义的短行被误判为可丢弃：{t}"
    for t in ("—", "***", "   ", "", None):
        assert not ok(t), f"纯标点/装饰符不应被当正文：{t}"


def test_dup_phrase_dedup_keeps_section_numbers():
    """叠词去重正则不得吃掉 `X.Y.Z` 编号（旧写法会把 4.2.2 收敛成 4.2）。"""
    import re
    pat = re.compile(r'([一-龥][一-龥A-Za-z]{1,9})\1')
    dedup = lambda s: pat.sub(lambda m: m.group(1), s)
    assert dedup("4.2.2 合理性评估") == "4.2.2 合理性评估"
    assert dedup("朱坝街道三圩社区朱坝街道三圩社区") == "朱坝街道三圩社区"
    import inspect
    from app.services.report_assembler import ReportAssembler
    src = inspect.getsource(ReportAssembler._post_generation_fix)
    assert "[一-龥][一-龥A-Za-z]{1,9}" in src, "叠词正则必须限定为「中文开头」"


def test_merge_first_col_merges_repeats_only():
    """首列相邻同名纵向合并（业主：「部分表格的第一列要合并同类项」）。

    三条边界必须守住：① 同名才并；② 合并格文字只留一行（不能叠成多行）；
    ③「合计」行绝不并入上一行。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from app.services.report_assembler import ReportAssembler

    d = Document()
    t = d.add_table(rows=6, cols=3, style='Table Grid')
    for ri, row in enumerate([("调查内容", "选项", "人数"),
                              ("您的身份", "本地居民", "48"),
                              ("您的身份", "租住本区", "1"),
                              ("您的年龄", "16～35", "5"),
                              ("您的年龄", "36～55", "26"),
                              ("合计", "—", "50")]):
        for ci, v in enumerate(row):
            t.cell(ri, ci).text = v

    assert ReportAssembler._merge_first_col(t) == 2, "应合并「您的身份」「您的年龄」两组"

    def merged(r):
        tcPr = t.rows[r].cells[0]._tc.find(qn('w:tcPr'))
        return tcPr is not None and tcPr.find(qn('w:vMerge')) is not None

    assert merged(1) and merged(3), "同名组首格应带 vMerge"
    assert t.rows[1].cells[0].text.strip() == "您的身份", "合并格文字不该叠成多行"
    assert not merged(5), "「合计」行不得并入上一行"
